"""
Tests for data-file format detection (purge_sim/data/formats.py), the point-by-point
(PxP) pressure sheet parser (data/pxp_parser.py) and the fast xlsx reader. The workbooks
are small synthetic copies of the real layouts, built here; no client data is checked in.
"""
import os
import sys

import openpyxl
import pytest

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, REPO_ROOT)

from purge_sim.app.importers import scenario_from_file   # noqa: E402
from purge_sim.app.workspace import Workspace             # noqa: E402
from purge_sim.data.formats import detect_format          # noqa: E402
from purge_sim.data.pxp_parser import parse_pxp           # noqa: E402
from purge_sim.data.xlsx_reader import peek_xlsx          # noqa: E402

# The PxP short field names (header row 4), in the sheet's column order up to PumpLoc.
PXP_FIELDS = ['LineID', 'Route', 'Series', 'Station', 'Measure', 'x', 'y', 'z', 'PipeOD', 'PipeWT', 'SMYS',
              'Mill', 'Seam', 'Year', 'VlaveType', 'MOV', 'ANSI', 'Descriptions', 'Feature', 'HPA', 'Dist',
              'MilePost', 'Elevation', 'Volume', 'User', 'MOPLimit', 'MOP', 'MOPvsPSMYS', 'Pop', 'PopvsPSMYS',
              'PSMYS', 'PDesign', 'Ptest', 'PTestvsPSMYS', 'MOPLimitSMYS', 'MOPLimitHT', 'MOPLimitHeadLight',
              'MOPLimitHeadHeavy', 'MOPCtrlHeadLight', 'MOPCtrlHeadHeavy', 'MOPCtrlLight', 'MOPCtrlHeavy',
              'MOPHeadLight', 'MOPHeadHeavy', 'PoCtrlHeadLight', 'PoCtrlHeadHeavy', 'PoCtrlLight', 'PoCtrlHeavy',
              'PoHeadLight', 'PoHeadHeavy', 'FlangeLoc', 'PumpLoc']
C = {name: i for i, name in reversed(list(enumerate(PXP_FIELDS)))}


def _pxp_row(dist, elev, mop, desc=None, feature="COORD_PT", valve=None, pump=False, wt=0.281):
    r = [None] * len(PXP_FIELDS)
    r[C['Dist']], r[C['Elevation']], r[C['MOP']], r[C['MOPLimit']] = dist, elev, mop, mop + 200
    r[C['PipeOD']], r[C['PipeWT']], r[C['Descriptions']], r[C['Feature']] = 12.75, wt, desc, feature
    r[C['VlaveType']] = valve
    r[C['MilePost']] = 300 - dist          # station mileposts run backwards on the real sheet
    r[C['PumpLoc']] = elev if pump else "#N/A"
    r[C['y']], r[C['x']] = 46.0 + (dist - 100) * 0.0145, -108.0   # due north, ~1 mi per mi
    return r


def write_pxp(path):
    wb = openpyxl.Workbook()
    guide = wb.active
    guide.title = "GuidePxPSheet"
    guide.append([None, "Point by Point (PxP) MOP and Operating Pressure Sheet"])
    ws = wb.create_sheet("PxPData")
    banner = [None] * len(PXP_FIELDS)
    banner[0], banner[2] = "SYSTEM NAME", "TEST 12 INCH CRUDE LINE"
    banner[15], banner[17], banner[18], banner[19] = "Specified Gravity - Light", 0.86, "Specified Gravity - Heavy", 0.92
    ws.append(banner)
    long1 = [None] * len(PXP_FIELDS)
    long1[C['Dist']], long1[C['Elevation']], long1[C['MOP']] = "Dist. from Origin (miles) ", "Elevation (ft)", "MOP (psig)"
    ws.append(long1)
    ws.append([None] * len(PXP_FIELDS))
    ws.append(PXP_FIELDS)
    for _ in range(3):
        ws.append([None] * len(PXP_FIELDS))
    rows = [
        _pxp_row(100.0, 4000, 1500, "PIPE", None),
        _pxp_row(100.01, 4001, 1500, "AGR", "AGR |  ORIGIN CHECK VALVE"),
        _pxp_row(100.02, 4001, 1560, "PMP", "ORIGIN STATION", pump=True),
    ]
    dist, elev = 100.03, 4001.0
    while dist < 120:
        dist = round(dist + 0.05, 3)
        elev += 3.0 if dist < 110 else -4.0
        rows.append(_pxp_row(dist, elev, 1500 - (elev - 4000) * 0.37))
    rows += [
        _pxp_row(104.0, 4030, 1489, "CHECK", "G12 104.0A - ELLIOT CV", valve="CHECK"),
        _pxp_row(112.0, 4050, 1481, "PMP", "MIDWAY STATION", pump=True, wt=0.375),
        _pxp_row(112.01, 4050, 1481, "CHECK", "MIDWAY CHECK VALVE", valve="CHECK", wt=0.375),
        _pxp_row(112.02, 4050, 1481, "BALL", "MIDWAY OUTSTN", valve="BALL", wt=0.375),
        _pxp_row(120.05, 3960, 1400, "REC", "END RECEIVER"),
    ]
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def write_rosen(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(9):
        ws.append(["Rosen report"])
    ws.append(["Calculated Mile Post [mi.]", "Elevation [ft.]", "Pipe Diameter [in.]",
               "Pipe Nominal Wall Thickness [in.]", "Maximum Operating Pressure (MOP) [psi]", "Site",
               "Additional description"])
    ws.append(["[mi.]", "[ft.]", "[in.]", "[in.]", "[psi]", None, None])
    for i in range(50):
        desc = "station check valve" if i == 20 else ("check valve" if i == 35 else None)
        ws.append([i * 0.1, 500 + i, 24.0, 0.313, 1440, "Site A" if i == 20 else None, desc])
    wb.save(path)
    return path


def write_plain_profile(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Milepost", "Elevation"])
    for i in range(20):
        ws.append([i * 0.5, 1000 + 10 * i])
    wb.save(path)
    return path


# ---------------------------------------------------------------- detection

def test_detect_format(tmp_path):
    assert detect_format(write_pxp(str(tmp_path / "line_pxp.xlsx"))) == "pxp"
    assert detect_format(write_rosen(str(tmp_path / "rosen.xlsx"))) == "ili"
    assert detect_format(write_plain_profile(str(tmp_path / "profile.xlsx"))) == "profile"
    for name in ("route.kmz", "route.kml", "gps.txt", "points.csv"):
        assert detect_format(str(tmp_path / name)) == "profile"     # by extension, not opened


def test_xlsx_reader_matches_openpyxl(tmp_path):
    path = write_pxp(str(tmp_path / "pxp.xlsx"))
    fast = peek_xlsx(path, max_rows=None, sheet="PxPData")["PxPData"]
    ws = openpyxl.load_workbook(path, data_only=True)["PxPData"]
    slow = [list(r) for r in ws.iter_rows(values_only=True)]
    assert len(fast) == len(slow)
    for a, b in zip(fast, slow):
        a = a + [None] * (len(b) - len(a))
        for x, y in zip(a, b):
            if isinstance(y, (int, float)):
                assert float(x) == pytest.approx(y)
            else:
                assert x == y


# ---------------------------------------------------------------- PxP parser

def test_parse_pxp(tmp_path):
    d = parse_pxp(write_pxp(str(tmp_path / "pxp.xlsx")))
    assert d.format == "PxP" and d.system_name == "TEST 12 INCH CRUDE LINE"
    assert (d.sg_light, d.sg_heavy) == (0.86, 0.92)
    prof = d.elevation_profile
    assert prof[0, 0] == pytest.approx(100.0) and prof[-1, 0] == pytest.approx(120.05)
    assert all(b > a for a, b in zip(prof[:-1, 0], prof[1:, 0]))      # sorted, no duplicate mileposts
    # the station at the launch end is the purge origin, not a pump ahead of the pig
    assert [s.name for s in d.pump_station_records] == ["Midway Station"]
    assert d.pump_station_records[0].check_valve_mp == pytest.approx(112.01)
    assert any("Origin Station" in n for n in d.notes)
    cvs = d.check_valves()
    assert [(round(c.mp, 2), c.is_pump_station) for c in cvs] == [(104.0, False), (112.01, True)]
    assert d.bpcv_record is None
    # per-point MOP is the sheet's MOP column, not the higher MOP Limit
    assert max(j.mop_psig for j in d.mop_joints) == pytest.approx(d.raw_df["mop"].max())
    assert max(j.mop_psig for j in d.mop_joints) < d.raw_df["mop_limit"].max()
    assert {wt for _, _, _, wt in d.pipe_geometry.segments} == {0.281, 0.375}


def test_import_auto_detects_pxp(tmp_path):
    path = write_pxp(str(tmp_path / "pxp.xlsx"))
    sc, notes = scenario_from_file(path, "auto", "pxp job")
    inp = sc.inputs
    assert inp.data_source == "PxP" and sc.meta.intake["format"] == "pxp"
    assert inp.purge_start_mp == pytest.approx(100.0) and inp.purge_end_mp == pytest.approx(120.05)
    assert [p["name"] for p in inp.pump_stations] == ["Midway Station"]
    assert len(inp.check_valves) == 2 and inp.mop_joints
    assert "point-by-point" in sc.meta.notes and "TEST 12 INCH CRUDE LINE" in sc.meta.notes
    # the sheet's x/y columns become the route for the Map tab
    assert len(inp.route_latlon) > 2 and inp.route_latlon[0] == pytest.approx([46.0, -108.0])
    # the pipe size came from the file, so the job setup doesn't ask for it
    ws = Workspace()
    r = ws.new_from_import(path, "auto", "pxp job")
    assert r["format"] == "pxp" and "PxP" in r["format_label"]
    assert "nps" not in r["job_setup"]["missing"]


def test_import_auto_detects_rosen(tmp_path):
    sc, _ = scenario_from_file(write_rosen(str(tmp_path / "rosen.xlsx")), "auto", "ili job")
    assert sc.inputs.data_source == "ILI" and sc.meta.intake["format"] == "ili"
    assert [p["name"] for p in sc.inputs.pump_stations] == ["Site A"]


def test_import_rejects_unknown_kind(tmp_path):
    with pytest.raises(ValueError):
        scenario_from_file(write_rosen(str(tmp_path / "rosen.xlsx")), "pdf", "x")
